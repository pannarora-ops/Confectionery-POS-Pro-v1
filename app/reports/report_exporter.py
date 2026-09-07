"""
Excel Report Exporter.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font


class ReportExporter:
    """Export reports to Excel."""

    def __init__(self):
        pass

    # ---------------------------------------------------------
    # Common
    # ---------------------------------------------------------

    def _create_sheet(
        self,
        title: str,
        headers: list[str],
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = title

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(
                row=1,
                column=col,
            )
            cell.value = header
            cell.font = Font(bold=True)

        return wb, ws

    def _save(
        self,
        workbook: Workbook,
        output_file: str,
    ) -> str:

        path = Path(output_file)

        workbook.save(path)

        return str(path)

    # ---------------------------------------------------------
    # Inventory
    # ---------------------------------------------------------

    def export_inventory(
        self,
        inventory: list[dict],
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Inventory",
            [
                "ID",
                "SKU",
                "Product",
                "Stock",
                "Purchase Price",
                "Selling Price",
                "Stock Value",
            ],
        )

        row = 2

        for product in inventory:

            ws.cell(row, 1).value = product["id"]
            ws.cell(row, 2).value = product["sku"]
            ws.cell(row, 3).value = product["name"]
            ws.cell(row, 4).value = float(product["stock"])
            ws.cell(row, 5).value = float(
                product["purchase_price"]
            )
            ws.cell(row, 6).value = float(
                product["selling_price"]
            )
            ws.cell(row, 7).value = float(
                product["stock_value"]
            )

            row += 1

        return self._save(
            wb,
            output_file,
        )

    # ---------------------------------------------------------
    # Sales
    # ---------------------------------------------------------

    def export_sales(
        self,
        sales: Iterable,
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Sales",
            [
                "Invoice",
                "Date",
                "Customer",
                "Subtotal",
                "GST",
                "Grand Total",
            ],
        )

        row = 2

        for sale in sales:

            ws.cell(row, 1).value = sale.invoice_number
            ws.cell(row, 2).value = str(
                sale.sale_date
            )
            ws.cell(row, 3).value = (
                sale.customer_name
                or (
                    sale.customer.name
                    if sale.customer
                    else ""
                )
            )
            ws.cell(row, 4).value = float(
                sale.subtotal
            )
            ws.cell(row, 5).value = float(
                sale.gst_total
            )
            ws.cell(row, 6).value = float(
                sale.grand_total
            )

            row += 1

        return self._save(
            wb,
            output_file,
        )

    # ---------------------------------------------------------
    # Purchases
    # ---------------------------------------------------------

    def export_purchases(
        self,
        purchases: Iterable,
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Purchases",
            [
                "Invoice",
                "Date",
                "Supplier",
                "Subtotal",
                "GST",
                "Grand Total",
            ],
        )

        row = 2

        for purchase in purchases:

            ws.cell(row, 1).value = (
                purchase.invoice_number
            )

            ws.cell(row, 2).value = str(
                purchase.purchase_date
            )

            ws.cell(row, 3).value = (
                purchase.supplier.name
                if purchase.supplier
                else ""
            )

            ws.cell(row, 4).value = float(
                purchase.subtotal
            )

            ws.cell(row, 5).value = float(
                purchase.gst_total
            )

            ws.cell(row, 6).value = float(
                purchase.grand_total
            )

            row += 1

        return self._save(
            wb,
            output_file,
        )

    # ---------------------------------------------------------
    # Customer Ledger
    # ---------------------------------------------------------

    def export_customer_ledger(
        self,
        entries: Iterable,
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Customer Ledger",
            [
                "Customer",
                "Type",
                "Reference",
                "Debit",
                "Credit",
                "Balance",
            ],
        )

        row = 2

        for item in entries:

            ws.cell(row, 1).value = (
                item.customer.name
                if item.customer
                else ""
            )

            ws.cell(row, 2).value = (
                item.transaction_type.value
            )

            ws.cell(row, 3).value = (
                item.reference_number
            )

            ws.cell(row, 4).value = float(
                item.debit
            )

            ws.cell(row, 5).value = float(
                item.credit
            )

            ws.cell(row, 6).value = float(
                item.balance
            )

            row += 1

        return self._save(
            wb,
            output_file,
        )

    # ---------------------------------------------------------
    # Supplier Ledger
    # ---------------------------------------------------------

    def export_supplier_ledger(
        self,
        entries: Iterable,
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Supplier Ledger",
            [
                "Supplier",
                "Type",
                "Reference",
                "Debit",
                "Credit",
                "Balance",
            ],
        )

        row = 2

        for item in entries:

            ws.cell(row, 1).value = (
                item.supplier.name
                if item.supplier
                else ""
            )

            ws.cell(row, 2).value = (
                item.transaction_type.value
            )

            ws.cell(row, 3).value = (
                item.reference_number
            )

            ws.cell(row, 4).value = float(
                item.debit
            )

            ws.cell(row, 5).value = float(
                item.credit
            )

            ws.cell(row, 6).value = float(
                item.balance
            )

            row += 1

        return self._save(
            wb,
            output_file,
        )

    # ---------------------------------------------------------
    # Stock History
    # ---------------------------------------------------------

    def export_stock_history(
        self,
        transactions: Iterable,
        output_file: str,
    ) -> str:

        wb, ws = self._create_sheet(
            "Stock History",
            [
                "Product",
                "Type",
                "Quantity",
                "Reference",
                "Remarks",
            ],
        )

        row = 2

        for tx in transactions:

            ws.cell(row, 1).value = (
                tx.product.name
                if tx.product
                else ""
            )

            ws.cell(row, 2).value = (
                tx.transaction_type.value
            )

            ws.cell(row, 3).value = float(
                tx.quantity
            )

            ws.cell(row, 4).value = tx.reference

            ws.cell(row, 5).value = tx.remarks

            row += 1

        return self._save(
            wb,
            output_file,
        )
    